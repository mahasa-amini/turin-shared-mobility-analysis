from pathlib import Path
import math
import re
import sys

import numpy as np
import pandas as pd
import geopandas as gpd

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
except ImportError:
    raise SystemExit(
        "\npython-docx is required for Appendix A generation.\n"
        "Install it once with:\n\n"
        "    pip install python-docx\n"
    )


# =============================================================================
# 1. PROJECT PATHS
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

OD_TABLE_DIR = (
    PROJECT_ROOT
    / "outputs"
    / "tables"
    / "exercise_2"
    / "od"
)

GIS_GPKG = (
    PROJECT_ROOT
    / "data"
    / "processed"
    / "exercise_2"
    / "gis"
    / "exercise2_od_layers.gpkg"
)

OUTPUT_DIR = (
    PROJECT_ROOT
    / "outputs"
    / "report"
    / "exercise_2"
    / "appendices"
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

OUTPUT_DOCX = (
    OUTPUT_DIR
    / "Appendix_A_complete_2024_all_day_OD_matrix.docx"
)

OUTPUT_MATRIX_CSV = (
    OUTPUT_DIR
    / "Appendix_A_system_all_day_OD_matrix_94x94.csv"
)

OUTPUT_VALIDATION_CSV = (
    OUTPUT_DIR
    / "Appendix_A_validation.csv"
)


# =============================================================================
# 2. EXPECTED VALUES FROM THE VALIDATED OD PIPELINE
# =============================================================================

EXPECTED_ZONE_COUNT = 94

EXPECTED_TOTAL_TRIPS = 1_681_927
EXPECTED_INTRA_TRIPS = 272_222
EXPECTED_INTER_TRIPS = 1_409_705

EXPECTED_ACTIVE_OD_PAIRS = 6_154
EXPECTED_ACTIVE_INTERZONAL_PAIRS = 6_070

# Maximum destination columns shown in each Appendix block.
#
# 94 destinations / 14 per block = 7 blocks.
DESTINATIONS_PER_BLOCK = 14


# =============================================================================
# 3. GENERAL UTILITIES
# =============================================================================

def normalize_zone_id(value):
    """
    Convert zone identifiers to the canonical form used in Exercise 2.

    Examples:
        1       -> 01
        "1"     -> 01
        "01"    -> 01
        9bis    -> 09bis
        "09bis" -> 09bis
        17bis   -> 17bis
    """

    if pd.isna(value):
        return None

    text = str(value).strip()

    if not text:
        return None

    text = text.replace(" ", "")

    # Numeric values accidentally represented as floats.
    numeric_match = re.fullmatch(
        r"(\d+)(?:\.0+)?",
        text,
    )

    if numeric_match:
        number = int(
            numeric_match.group(1)
        )

        return f"{number:02d}"

    # Alphanumeric official IDs such as 09bis and 17bis.
    bis_match = re.fullmatch(
        r"(\d+)(bis)",
        text,
        flags=re.IGNORECASE,
    )

    if bis_match:
        number = int(
            bis_match.group(1)
        )

        return f"{number:02d}bis"

    return text


def zone_sort_key(zone_id):
    """
    Sort official IDs naturally:
        01, 02, ..., 09, 09bis, 10, ..., 17, 17bis, ...
    """

    zone_id = normalize_zone_id(
        zone_id
    )

    match = re.fullmatch(
        r"(\d+)(.*)",
        zone_id,
    )

    if match is None:
        return (
            9999,
            zone_id,
        )

    number = int(
        match.group(1)
    )

    suffix = (
        match.group(2)
        .lower()
        .strip()
    )

    # Empty suffix comes before "bis".
    suffix_rank = (
        0
        if suffix == ""
        else 1
    )

    return (
        number,
        suffix_rank,
        suffix,
    )


def print_header(title):
    print(
        "\n"
        + "=" * 110
    )
    print(title)
    print(
        "=" * 110
    )


# =============================================================================
# 4. LOAD OFFICIAL ZONES
# =============================================================================

def load_official_zones():
    """
    Load zones_all_day so that:
    - official zone IDs define matrix ordering;
    - production totals validate matrix rows;
    - attraction totals validate matrix columns.
    """

    if not GIS_GPKG.exists():
        raise FileNotFoundError(
            f"GIS GeoPackage not found:\n{GIS_GPKG}"
        )

    zones = gpd.read_file(
        GIS_GPKG,
        layer="zones_all_day",
        engine="pyogrio",
    )

    required = {
        "zone_id",
        "zone_name",
        "production_trips",
        "attraction_trips",
    }

    missing = (
        required
        - set(zones.columns)
    )

    if missing:
        raise KeyError(
            "zones_all_day is missing required columns:\n"
            f"{sorted(missing)}"
        )

    zones = zones[
        [
            "zone_id",
            "zone_name",
            "production_trips",
            "attraction_trips",
        ]
    ].copy()

    zones["zone_id"] = (
        zones["zone_id"]
        .map(normalize_zone_id)
    )

    if zones["zone_id"].duplicated().any():
        duplicates = (
            zones.loc[
                zones["zone_id"].duplicated(
                    keep=False
                ),
                "zone_id",
            ]
            .tolist()
        )

        raise ValueError(
            "Duplicate official zone IDs detected:\n"
            f"{duplicates}"
        )

    zones = (
        zones
        .sort_values(
            "zone_id",
            key=lambda s: s.map(
                zone_sort_key
            ),
        )
        .reset_index(
            drop=True
        )
    )

    if len(zones) != EXPECTED_ZONE_COUNT:
        raise ValueError(
            f"Expected {EXPECTED_ZONE_COUNT} zones, "
            f"found {len(zones)}."
        )

    return zones


# =============================================================================
# 5. SEARCH FOR EXISTING SQUARE SYSTEM MATRIX
# =============================================================================

def try_read_square_matrix(path, official_zone_ids):
    """
    Test whether a CSV contains the complete 94 x 94 system matrix.

    Expected general structure:

        origin_zone_id, 01, 02, 03, ..., 94
        01, ...
        02, ...
        ...

    The first column name may vary.
    """

    try:
        df = pd.read_csv(
            path,
            dtype=str,
        )
    except Exception:
        return None

    if len(df) != EXPECTED_ZONE_COUNT:
        return None

    # Identify destination columns from their headers.
    destination_columns = {}

    for col in df.columns:

        normalized = normalize_zone_id(
            col
        )

        if normalized in official_zone_ids:
            destination_columns[
                normalized
            ] = col

    if len(destination_columns) != EXPECTED_ZONE_COUNT:
        return None

    # Identify the origin-ID column.
    non_destination_cols = [
        col
        for col in df.columns
        if col
        not in destination_columns.values()
    ]

    origin_col = None

    # Prefer a sensible named field.
    preferred_origin_names = {
        "origin_zone_id",
        "origin_zone",
        "zone_id",
        "origin",
        "unnamed: 0",
    }

    for col in non_destination_cols:

        if str(col).lower() in preferred_origin_names:
            origin_col = col
            break

    # Otherwise find a column containing all 94 official IDs.
    if origin_col is None:

        for col in non_destination_cols:

            normalized_values = (
                df[col]
                .map(normalize_zone_id)
            )

            if (
                set(normalized_values.dropna())
                == set(official_zone_ids)
            ):
                origin_col = col
                break

    if origin_col is None:
        return None

    origins = (
        df[origin_col]
        .map(normalize_zone_id)
    )

    if set(origins) != set(official_zone_ids):
        return None

    matrix = pd.DataFrame(
        index=origins,
        columns=official_zone_ids,
        dtype=float,
    )

    for destination_id in official_zone_ids:

        source_column = (
            destination_columns[
                destination_id
            ]
        )

        values = pd.to_numeric(
            df[source_column],
            errors="coerce",
        )

        if values.isna().any():
            return None

        matrix.loc[
            origins,
            destination_id,
        ] = values.to_numpy()

    matrix = (
        matrix
        .reindex(
            index=official_zone_ids,
            columns=official_zone_ids,
        )
    )

    # Complete OD matrix must contain integers.
    if not np.allclose(
        matrix.to_numpy(),
        np.round(
            matrix.to_numpy()
        ),
    ):
        return None

    matrix = (
        matrix
        .round()
        .astype("int64")
    )

    if (
        int(
            matrix.to_numpy().sum()
        )
        != EXPECTED_TOTAL_TRIPS
    ):
        return None

    return matrix


def find_existing_square_matrix(
    official_zone_ids,
):
    """
    Search existing Exercise 2 OD output CSVs for the validated
    94 x 94 system matrix.
    """

    if not OD_TABLE_DIR.exists():
        raise FileNotFoundError(
            f"OD table directory not found:\n"
            f"{OD_TABLE_DIR}"
        )

    csv_files = sorted(
        OD_TABLE_DIR.rglob(
            "*.csv"
        )
    )

    if not csv_files:
        raise FileNotFoundError(
            "No CSV files were found in:\n"
            f"{OD_TABLE_DIR}"
        )

    print(
        f"Searching {len(csv_files)} existing OD CSV files..."
    )

    # Prioritize filenames that are likely to be system square matrices.
    def candidate_priority(path):

        name = (
            path.name
            .lower()
        )

        score = 0

        if "system" in name:
            score -= 10

        if "matrix" in name:
            score -= 8

        if "square" in name:
            score -= 5

        if "all_day" in name or "allday" in name:
            score -= 4

        if "2024" in name:
            score -= 2

        return (
            score,
            name,
        )

    csv_files = sorted(
        csv_files,
        key=candidate_priority,
    )

    for path in csv_files:

        matrix = (
            try_read_square_matrix(
                path,
                official_zone_ids,
            )
        )

        if matrix is not None:
            return (
                matrix,
                path,
                "existing_square_matrix",
            )

    return None


# =============================================================================
# 6. FALLBACK: RECONSTRUCT MATRIX FROM A LONG-FORM SYSTEM OD CSV
# =============================================================================

ORIGIN_CANDIDATES = [
    "origin_zone_id",
    "origin_zone",
    "origin",
    "o_zone_id",
    "o_zone",
    "from_zone_id",
    "from_zone",
]

DESTINATION_CANDIDATES = [
    "destination_zone_id",
    "destination_zone",
    "destination",
    "dest_zone_id",
    "dest_zone",
    "d_zone_id",
    "d_zone",
    "to_zone_id",
    "to_zone",
]

COUNT_CANDIDATES = [
    "system_trips",
    "trips",
    "trip_count",
    "od_trips",
    "flow_trips",
    "flow",
    "count",
]


def find_column(
    columns,
    candidates,
):
    lower_map = {
        str(col).lower(): col
        for col in columns
    }

    for candidate in candidates:

        if candidate in lower_map:
            return lower_map[
                candidate
            ]

    return None


def try_read_long_matrix(
    path,
    official_zone_ids,
):
    """
    Try to reconstruct the system matrix from a long-format OD table.
    """

    try:
        df = pd.read_csv(
            path,
            low_memory=False,
        )
    except Exception:
        return None

    origin_col = find_column(
        df.columns,
        ORIGIN_CANDIDATES,
    )

    destination_col = find_column(
        df.columns,
        DESTINATION_CANDIDATES,
    )

    count_col = find_column(
        df.columns,
        COUNT_CANDIDATES,
    )

    if (
        origin_col is None
        or destination_col is None
        or count_col is None
    ):
        return None

    temp = df[
        [
            origin_col,
            destination_col,
            count_col,
        ]
    ].copy()

    temp["origin_zone_id"] = (
        temp[origin_col]
        .map(normalize_zone_id)
    )

    temp["destination_zone_id"] = (
        temp[destination_col]
        .map(normalize_zone_id)
    )

    temp["trips"] = pd.to_numeric(
        temp[count_col],
        errors="coerce",
    )

    temp = temp.dropna(
        subset=[
            "origin_zone_id",
            "destination_zone_id",
            "trips",
        ]
    )

    temp = temp[
        temp["origin_zone_id"].isin(
            official_zone_ids
        )
        & temp[
            "destination_zone_id"
        ].isin(
            official_zone_ids
        )
    ].copy()

    if temp.empty:
        return None

    grouped = (
        temp
        .groupby(
            [
                "origin_zone_id",
                "destination_zone_id",
            ],
            as_index=False,
        )["trips"]
        .sum()
    )

    matrix = (
        grouped
        .pivot(
            index="origin_zone_id",
            columns="destination_zone_id",
            values="trips",
        )
        .reindex(
            index=official_zone_ids,
            columns=official_zone_ids,
            fill_value=0,
        )
        .fillna(0)
    )

    if not np.allclose(
        matrix.to_numpy(),
        np.round(
            matrix.to_numpy()
        ),
    ):
        return None

    matrix = (
        matrix
        .round()
        .astype("int64")
    )

    if (
        int(
            matrix.to_numpy().sum()
        )
        != EXPECTED_TOTAL_TRIPS
    ):
        return None

    return matrix


def find_long_format_matrix(
    official_zone_ids,
):
    """
    Fallback if the square matrix CSV cannot be found.
    """

    csv_files = sorted(
        OD_TABLE_DIR.rglob(
            "*.csv"
        )
    )

    def priority(path):

        name = (
            path.name
            .lower()
        )

        score = 0

        if "system" in name:
            score -= 10

        if "long" in name:
            score -= 8

        if "od" in name:
            score -= 4

        if "all_day" in name:
            score -= 3

        return (
            score,
            name,
        )

    csv_files = sorted(
        csv_files,
        key=priority,
    )

    for path in csv_files:

        matrix = (
            try_read_long_matrix(
                path,
                official_zone_ids,
            )
        )

        if matrix is not None:
            return (
                matrix,
                path,
                "reconstructed_from_long_OD",
            )

    return None


# =============================================================================
# 7. LOAD MATRIX
# =============================================================================

def load_system_matrix(
    official_zone_ids,
):
    """
    First try the existing 94 x 94 matrix.
    If unavailable, reconstruct it from the validated
    long-format OD output.
    """

    result = (
        find_existing_square_matrix(
            official_zone_ids
        )
    )

    if result is not None:
        return result

    print(
        "No validated square matrix CSV found."
    )
    print(
        "Trying long-format system OD outputs..."
    )

    result = (
        find_long_format_matrix(
            official_zone_ids
        )
    )

    if result is not None:
        return result

    raise FileNotFoundError(
        "\nCould not locate a validated all-day system OD matrix "
        "inside:\n"
        f"{OD_TABLE_DIR}\n\n"
        "Expected either:\n"
        "- a 94 x 94 square system matrix CSV, or\n"
        "- a long-format OD CSV with origin, destination and trip-count fields.\n"
    )


# =============================================================================
# 8. MATRIX VALIDATION
# =============================================================================

def validate_matrix(
    matrix,
    zones,
):
    """
    Perform independent Appendix-level validation.
    """

    zone_ids = (
        zones["zone_id"]
        .tolist()
    )

    if matrix.shape != (
        EXPECTED_ZONE_COUNT,
        EXPECTED_ZONE_COUNT,
    ):
        raise AssertionError(
            f"Unexpected matrix shape: {matrix.shape}"
        )

    if list(matrix.index) != zone_ids:
        raise AssertionError(
            "Matrix origin order does not match official zone order."
        )

    if list(matrix.columns) != zone_ids:
        raise AssertionError(
            "Matrix destination order does not match official zone order."
        )

    values = (
        matrix
        .to_numpy(
            dtype=np.int64
        )
    )

    total = int(
        values.sum()
    )

    intra = int(
        np.trace(
            values
        )
    )

    inter = (
        total
        - intra
    )

    active_all = int(
        np.count_nonzero(
            values
        )
    )

    off_diagonal = (
        values.copy()
    )

    np.fill_diagonal(
        off_diagonal,
        0,
    )

    active_inter = int(
        np.count_nonzero(
            off_diagonal
        )
    )

    if total != EXPECTED_TOTAL_TRIPS:
        raise AssertionError(
            f"Total trips mismatch: {total:,}"
        )

    if intra != EXPECTED_INTRA_TRIPS:
        raise AssertionError(
            f"Intra-zonal trips mismatch: {intra:,}"
        )

    if inter != EXPECTED_INTER_TRIPS:
        raise AssertionError(
            f"Inter-zonal trips mismatch: {inter:,}"
        )

    if active_all != EXPECTED_ACTIVE_OD_PAIRS:
        raise AssertionError(
            f"Active OD pairs mismatch: {active_all:,}"
        )

    if (
        active_inter
        != EXPECTED_ACTIVE_INTERZONAL_PAIRS
    ):
        raise AssertionError(
            "Active inter-zonal pair mismatch: "
            f"{active_inter:,}"
        )

    # ---------------------------------------------------------
    # Row totals = production
    # ---------------------------------------------------------

    matrix_production = (
        matrix.sum(
            axis=1
        )
        .astype("int64")
    )

    expected_production = (
        zones
        .set_index(
            "zone_id"
        )[
            "production_trips"
        ]
        .reindex(
            zone_ids
        )
        .astype("int64")
    )

    production_diff = (
        matrix_production
        - expected_production
    )

    if (
        production_diff
        != 0
    ).any():

        bad = (
            production_diff[
                production_diff != 0
            ]
        )

        raise AssertionError(
            "Matrix row totals do not reconcile with "
            "zones_all_day production_trips:\n"
            f"{bad}"
        )

    # ---------------------------------------------------------
    # Column totals = attraction
    # ---------------------------------------------------------

    matrix_attraction = (
        matrix.sum(
            axis=0
        )
        .astype("int64")
    )

    expected_attraction = (
        zones
        .set_index(
            "zone_id"
        )[
            "attraction_trips"
        ]
        .reindex(
            zone_ids
        )
        .astype("int64")
    )

    attraction_diff = (
        matrix_attraction
        - expected_attraction
    )

    if (
        attraction_diff
        != 0
    ).any():

        bad = (
            attraction_diff[
                attraction_diff != 0
            ]
        )

        raise AssertionError(
            "Matrix column totals do not reconcile with "
            "zones_all_day attraction_trips:\n"
            f"{bad}"
        )

    validation = pd.DataFrame(
        [
            {
                "metric": "official_zones",
                "value": EXPECTED_ZONE_COUNT,
                "status": "PASS",
            },
            {
                "metric": "matrix_rows",
                "value": matrix.shape[0],
                "status": "PASS",
            },
            {
                "metric": "matrix_columns",
                "value": matrix.shape[1],
                "status": "PASS",
            },
            {
                "metric": "possible_od_cells",
                "value": matrix.size,
                "status": "PASS",
            },
            {
                "metric": "total_valid_od_trips",
                "value": total,
                "status": "PASS",
            },
            {
                "metric": "intra_zonal_trips",
                "value": intra,
                "status": "PASS",
            },
            {
                "metric": "inter_zonal_trips",
                "value": inter,
                "status": "PASS",
            },
            {
                "metric": "active_od_pairs",
                "value": active_all,
                "status": "PASS",
            },
            {
                "metric": "active_interzonal_pairs",
                "value": active_inter,
                "status": "PASS",
            },
            {
                "metric": "row_totals_vs_production",
                "value": 0,
                "status": "PASS",
            },
            {
                "metric": "column_totals_vs_attraction",
                "value": 0,
                "status": "PASS",
            },
        ]
    )

    return validation


# =============================================================================
# 9. WORD / DOCX FORMATTING UTILITIES
# =============================================================================

def set_cell_shading(
    cell,
    fill,
):
    tc_pr = (
        cell._tc
        .get_or_add_tcPr()
    )

    shd = tc_pr.find(
        qn("w:shd")
    )

    if shd is None:
        shd = OxmlElement(
            "w:shd"
        )

        tc_pr.append(
            shd
        )

    shd.set(
        qn("w:fill"),
        fill,
    )


def set_repeat_table_header(
    row,
):
    tr_pr = (
        row._tr
        .get_or_add_trPr()
    )

    tbl_header = OxmlElement(
        "w:tblHeader"
    )

    tbl_header.set(
        qn("w:val"),
        "true",
    )

    tr_pr.append(
        tbl_header
    )


def prevent_row_split(
    row,
):
    tr_pr = (
        row._tr
        .get_or_add_trPr()
    )

    cant_split = OxmlElement(
        "w:cantSplit"
    )

    tr_pr.append(
        cant_split
    )


def set_row_height_exact(
    row,
    points,
):
    tr_pr = (
        row._tr
        .get_or_add_trPr()
    )

    tr_height = OxmlElement(
        "w:trHeight"
    )

    tr_height.set(
        qn("w:val"),
        str(
            int(
                points * 20
            )
        ),
    )

    tr_height.set(
        qn("w:hRule"),
        "exact",
    )

    tr_pr.append(
        tr_height
    )


def set_table_cell_margins(
    table,
    top=0,
    start=20,
    bottom=0,
    end=20,
):
    """
    Margins are specified in twentieths of a point (twips).
    """

    tbl = table._tbl

    tbl_pr = (
        tbl.tblPr
    )

    tbl_cell_mar = (
        tbl_pr.first_child_found_in(
            "w:tblCellMar"
        )
    )

    if tbl_cell_mar is None:

        tbl_cell_mar = OxmlElement(
            "w:tblCellMar"
        )

        tbl_pr.append(
            tbl_cell_mar
        )

    for margin_name, margin_value in [
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ]:

        node = (
            tbl_cell_mar.find(
                qn(
                    f"w:{margin_name}"
                )
            )
        )

        if node is None:

            node = OxmlElement(
                f"w:{margin_name}"
            )

            tbl_cell_mar.append(
                node
            )

        node.set(
            qn("w:w"),
            str(
                margin_value
            ),
        )

        node.set(
            qn("w:type"),
            "dxa",
        )


def set_cell_text(
    cell,
    text,
    font_size=5.5,
    bold=False,
    alignment=WD_ALIGN_PARAGRAPH.CENTER,
):
    cell.text = ""

    paragraph = (
        cell.paragraphs[0]
    )

    paragraph.alignment = (
        alignment
    )

    paragraph.paragraph_format.space_before = (
        Pt(0)
    )

    paragraph.paragraph_format.space_after = (
        Pt(0)
    )

    paragraph.paragraph_format.line_spacing = (
        Pt(font_size + 0.3)
    )

    run = paragraph.add_run(
        str(text)
    )

    run.font.name = (
        "Arial"
    )

    run.font.size = (
        Pt(font_size)
    )

    run.bold = (
        bold
    )

    cell.vertical_alignment = (
        WD_ALIGN_VERTICAL.CENTER
    )


# =============================================================================
# 10. CREATE APPENDIX WORD DOCUMENT
# =============================================================================

def build_appendix_docx(
    matrix,
    zones,
):
    zone_ids = (
        zones["zone_id"]
        .tolist()
    )

    doc = Document()

    # -------------------------------------------------------------------------
    # Page setup
    # -------------------------------------------------------------------------

    section = (
        doc.sections[0]
    )

    section.page_width = (
        Mm(210)
    )

    section.page_height = (
        Mm(297)
    )

    # Narrow but still printable margins.
    section.top_margin = (
        Mm(8)
    )

    section.bottom_margin = (
        Mm(8)
    )

    section.left_margin = (
        Mm(7)
    )

    section.right_margin = (
        Mm(7)
    )

    # -------------------------------------------------------------------------
    # Normal style
    # -------------------------------------------------------------------------

    normal_style = (
        doc.styles["Normal"]
    )

    normal_style.font.name = (
        "Arial"
    )

    normal_style.font.size = (
        Pt(9)
    )

    # -------------------------------------------------------------------------
    # Appendix title
    # -------------------------------------------------------------------------

    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    title.paragraph_format.space_after = (
        Pt(8)
    )

    run = title.add_run(
        "Appendix A — Complete 2024 System-wide "
        "Origin–Destination Matrix"
    )

    run.font.name = (
        "Arial"
    )

    run.font.size = (
        Pt(15)
    )

    run.bold = True

    # -------------------------------------------------------------------------
    # A.1
    # -------------------------------------------------------------------------

    heading = doc.add_paragraph()

    heading.paragraph_format.space_before = (
        Pt(4)
    )

    heading.paragraph_format.space_after = (
        Pt(5)
    )

    run = heading.add_run(
        "A.1 Matrix structure and interpretation"
    )

    run.font.name = (
        "Arial"
    )

    run.font.size = (
        Pt(11)
    )

    run.bold = True

    appendix_paragraphs = [
        (
            "The following tables report the complete system-wide "
            "all-day origin–destination matrix for the 94 official "
            "statistical zones of Turin. The matrix is based on all "
            "valid 2024 Bird, Lime and VOI trips for which both the "
            "origin and destination were assigned unambiguously to "
            "an official statistical zone."
        ),
        (
            "Rows represent origin zones and columns represent "
            "destination zones. Each cell reports the number of "
            "observed trips from the corresponding origin zone to "
            "the corresponding destination zone. Diagonal cells "
            "contain intra-zonal trips, while off-diagonal cells "
            "represent inter-zonal movements."
        ),
        (
            "The complete matrix contains 8,836 cells and sums to "
            "1,681,927 valid O–D trips. Zero-flow cells are retained "
            "explicitly. Because 94 destination columns cannot be "
            "displayed legibly on a standard report page, the matrix "
            "is presented in consecutive destination-zone blocks. "
            "All blocks retain the same 94 origin zones and together "
            "form the complete numerical matrix without aggregation "
            "or omission."
        ),
    ]

    for text in appendix_paragraphs:

        paragraph = (
            doc.add_paragraph()
        )

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )

        paragraph.paragraph_format.space_after = (
            Pt(5)
        )

        paragraph.paragraph_format.line_spacing = (
            1.05
        )

        run = paragraph.add_run(
            text
        )

        run.font.name = (
            "Arial"
        )

        run.font.size = (
            Pt(9)
        )

    note = doc.add_paragraph()

    note.paragraph_format.space_before = (
        Pt(6)
    )

    note.paragraph_format.space_after = (
        Pt(0)
    )

    run = note.add_run(
        "The seven table blocks presented in Section A.2 "
        "collectively form the complete 94 × 94 all-day O–D matrix."
    )

    run.font.name = (
        "Arial"
    )

    run.font.size = (
        Pt(9)
    )

    # -------------------------------------------------------------------------
    # Start matrix tables on a fresh page.
    # -------------------------------------------------------------------------

    doc.add_page_break()

    # -------------------------------------------------------------------------
    # Destination blocks
    # -------------------------------------------------------------------------

    blocks = [
        zone_ids[i : i + DESTINATIONS_PER_BLOCK]
        for i in range(
            0,
            len(zone_ids),
            DESTINATIONS_PER_BLOCK,
        )
    ]

    if len(blocks) != 7:
        raise AssertionError(
            f"Expected 7 blocks, obtained {len(blocks)}."
        )

    for block_number, destination_ids in enumerate(
        blocks,
        start=1,
    ):

        if block_number == 1:

            a2 = doc.add_paragraph()

            a2.paragraph_format.space_after = (
                Pt(2)
            )

            run = a2.add_run(
                "A.2 Complete 2024 system-wide all-day O–D matrix"
            )

            run.font.name = (
                "Arial"
            )

            run.font.size = (
                Pt(10)
            )

            run.bold = True

        # ---------------------------------------------------------------------
        # Table caption
        # ---------------------------------------------------------------------

        first_destination = (
            destination_ids[0]
        )

        last_destination = (
            destination_ids[-1]
        )

        caption = doc.add_paragraph()

        caption.paragraph_format.space_before = (
            Pt(0)
        )

        caption.paragraph_format.space_after = (
            Pt(2)
        )

        run = caption.add_run(
            f"Table A.{block_number}. "
            f"All-day O–D matrix — destination zones "
            f"{first_destination} to {last_destination}."
        )

        run.font.name = (
            "Arial"
        )

        run.font.size = (
            Pt(7)
        )

        # ---------------------------------------------------------------------
        # Create table
        # ---------------------------------------------------------------------

        table = doc.add_table(
            rows=1,
            cols=(
                1
                + len(
                    destination_ids
                )
            ),
        )

        table.style = (
            "Table Grid"
        )

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        table.autofit = False

        set_table_cell_margins(
            table,
            top=0,
            bottom=0,
            start=15,
            end=15,
        )

        # ---------------------------------------------------------------------
        # Width allocation
        # ---------------------------------------------------------------------

        usable_width_mm = (
            210
            - 7
            - 7
        )

        origin_width_mm = (
            12
        )

        destination_width_mm = (
            (
                usable_width_mm
                - origin_width_mm
            )
            / len(
                destination_ids
            )
        )

        # ---------------------------------------------------------------------
        # Header row
        # ---------------------------------------------------------------------

        header = (
            table.rows[0]
        )

        set_repeat_table_header(
            header
        )

        set_row_height_exact(
            header,
            7.2,
        )

        set_cell_text(
            header.cells[0],
            "O\\D",
            font_size=5.5,
            bold=True,
        )

        header.cells[0].width = (
            Mm(
                origin_width_mm
            )
        )

        set_cell_shading(
            header.cells[0],
            "D9D9D9",
        )

        for col_idx, destination_id in enumerate(
            destination_ids,
            start=1,
        ):

            cell = (
                header.cells[
                    col_idx
                ]
            )

            set_cell_text(
                cell,
                destination_id,
                font_size=5.5,
                bold=True,
            )

            cell.width = (
                Mm(
                    destination_width_mm
                )
            )

            set_cell_shading(
                cell,
                "D9D9D9",
            )

        # ---------------------------------------------------------------------
        # Matrix rows
        # ---------------------------------------------------------------------

        for origin_id in zone_ids:

            row = (
                table.add_row()
            )

            set_row_height_exact(
                row,
                6.7,
            )

            prevent_row_split(
                row
            )

            origin_cell = (
                row.cells[0]
            )

            set_cell_text(
                origin_cell,
                origin_id,
                font_size=5.3,
            )

            origin_cell.width = (
                Mm(
                    origin_width_mm
                )
            )

            set_cell_shading(
                origin_cell,
                "F2F2F2",
            )

            for col_idx, destination_id in enumerate(
                destination_ids,
                start=1,
            ):

                value = int(
                    matrix.loc[
                        origin_id,
                        destination_id,
                    ]
                )

                cell = (
                    row.cells[
                        col_idx
                    ]
                )

                set_cell_text(
                    cell,
                    value,
                    font_size=5.3,
                )

                cell.width = (
                    Mm(
                        destination_width_mm
                    )
                )

                # Lightly identify matrix diagonal.
                if (
                    origin_id
                    == destination_id
                ):
                    set_cell_shading(
                        cell,
                        "EAEAEA",
                    )

        # ---------------------------------------------------------------------
        # Footer note on each block
        # ---------------------------------------------------------------------

        footer = doc.add_paragraph()

        footer.paragraph_format.space_before = (
            Pt(1)
        )

        footer.paragraph_format.space_after = (
            Pt(0)
        )

        run = footer.add_run(
            "Rows = origin zones; columns = destination zones; "
            "entries = observed trip counts."
        )

        run.font.name = (
            "Arial"
        )

        run.font.size = (
            Pt(6)
        )

        # Start every subsequent block on its own page.
        if block_number < len(blocks):
            doc.add_page_break()

    doc.save(
        OUTPUT_DOCX
    )

    return len(blocks)


# =============================================================================
# 11. SAVE FULL MATRIX CSV
# =============================================================================

def save_matrix_csv(
    matrix,
):
    output = (
        matrix.copy()
    )

    output.index.name = (
        "origin_zone_id"
    )

    output.to_csv(
        OUTPUT_MATRIX_CSV,
        encoding="utf-8-sig",
    )


# =============================================================================
# 12. MAIN
# =============================================================================

def main():

    print_header(
        "EXERCISE 2 — APPENDIX A: COMPLETE ALL-DAY O–D MATRIX"
    )

    # -------------------------------------------------------------------------
    # Official zones
    # -------------------------------------------------------------------------

    print(
        "\nLoading official zoning..."
    )

    zones = (
        load_official_zones()
    )

    official_zone_ids = (
        zones["zone_id"]
        .tolist()
    )

    print(
        f"Official zones: {len(zones)}"
    )

    print(
        "First IDs:",
        official_zone_ids[:10],
    )

    print(
        "Last IDs:",
        official_zone_ids[-10:],
    )

    # -------------------------------------------------------------------------
    # Matrix
    # -------------------------------------------------------------------------

    print_header(
        "LOCATING EXISTING VALIDATED SYSTEM MATRIX"
    )

    matrix, source_path, source_type = (
        load_system_matrix(
            official_zone_ids
        )
    )

    print(
        f"Matrix source type: {source_type}"
    )

    print(
        f"Matrix source file:\n{source_path}"
    )

    print(
        f"Matrix shape: {matrix.shape}"
    )

    # -------------------------------------------------------------------------
    # Validation
    # -------------------------------------------------------------------------

    print_header(
        "APPENDIX A MATRIX VALIDATION"
    )

    validation = (
        validate_matrix(
            matrix,
            zones,
        )
    )

    print(
        validation.to_string(
            index=False
        )
    )

    validation.to_csv(
        OUTPUT_VALIDATION_CSV,
        index=False,
    )

    # -------------------------------------------------------------------------
    # Save canonical full matrix
    # -------------------------------------------------------------------------

    save_matrix_csv(
        matrix
    )

    # -------------------------------------------------------------------------
    # Build Word Appendix
    # -------------------------------------------------------------------------

    print_header(
        "BUILDING WORD APPENDIX"
    )

    block_count = (
        build_appendix_docx(
            matrix,
            zones,
        )
    )

    # -------------------------------------------------------------------------
    # Final summary
    # -------------------------------------------------------------------------

    print_header(
        "APPENDIX A COMPLETE"
    )

    total = int(
        matrix.to_numpy().sum()
    )

    intra = int(
        np.trace(
            matrix.to_numpy()
        )
    )

    inter = (
        total
        - intra
    )

    print(
        f"Matrix dimensions:       {matrix.shape[0]} x {matrix.shape[1]}"
    )

    print(
        f"Possible OD cells:       {matrix.size:,}"
    )

    print(
        f"Total valid OD trips:    {total:,}"
    )

    print(
        f"Intra-zonal trips:       {intra:,}"
    )

    print(
        f"Inter-zonal trips:       {inter:,}"
    )

    print(
        f"Destination blocks:      {block_count}"
    )

    print(
        "\nSaved full matrix CSV ->"
    )

    print(
        OUTPUT_MATRIX_CSV
    )

    print(
        "\nSaved validation table ->"
    )

    print(
        OUTPUT_VALIDATION_CSV
    )

    print(
        "\nSaved Word Appendix A ->"
    )

    print(
        OUTPUT_DOCX
    )

    print(
        "\nAll Appendix A validation checks: PASSED"
    )


if __name__ == "__main__":
    main()