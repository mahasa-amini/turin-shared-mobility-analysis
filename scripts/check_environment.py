"""Synthetic dependency checks; never import project analyses or read project data."""

import importlib.util
import os
from pathlib import Path
import tempfile
import unittest


class EnvironmentChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="transport-smoke-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_parquet(self):
        import pandas as pd

        expected = pd.DataFrame({
            "count": pd.array([3, None], dtype="Int64"),
            "value": [1.25, float("nan")],
            "label": pd.array(["Turin", None], dtype="string"),
            "time": pd.to_datetime(["2024-07-01T12:00:00Z", None], utc=True)
            .tz_convert("Europe/Rome"),
        })
        path = self.root / "sample.parquet"
        expected.to_parquet(path, engine="pyarrow", index=False)
        actual = pd.read_parquet(path, engine="pyarrow")
        pd.testing.assert_frame_equal(actual, expected)

    def test_excel(self):
        import pandas as pd

        expected = pd.DataFrame({"label": ["Turin", "Voi"], "value": [1.25, 2.5]})
        path = self.root / "sample.xlsx"
        expected.to_excel(path, engine="openpyxl", index=False)
        actual = pd.read_excel(path, engine="openpyxl")
        pd.testing.assert_frame_equal(actual, expected)

    def test_geopackage(self):
        import geopandas as gpd
        from shapely.geometry import Point

        expected = gpd.GeoDataFrame(
            {"id": [1, 2]}, geometry=[Point(0, 0), Point(3, 4)], crs="EPSG:3857"
        )
        path = self.root / "sample.gpkg"
        expected.to_file(path, layer="points", driver="GPKG", engine="pyogrio", index=False)
        actual = gpd.read_file(path, layer="points", engine="pyogrio").sort_values("id")
        self.assertEqual(actual.crs, expected.crs)
        self.assertEqual(actual["id"].tolist(), [1, 2])
        self.assertTrue(actual.geometry.geom_equals(expected.geometry).all())

    def test_projection_and_nearest_join(self):
        import geopandas as gpd
        from shapely.geometry import Point

        # One longitude degree at the equator in Web Mercator is R * pi / 180.
        projected = gpd.GeoSeries([Point(0, 0), Point(1, 0)], crs=4326).to_crs(3857)
        self.assertAlmostEqual(projected.iloc[1].x, 111319.49079327357, places=5)
        self.assertAlmostEqual(projected.iloc[1].y, 0.0, places=5)
        left = gpd.GeoDataFrame(geometry=[projected.iloc[0]], crs=projected.crs)
        right = gpd.GeoDataFrame(
            {"stop": ["near", "far"]},
            geometry=[Point(3, 4), Point(30, 40)], crs=projected.crs,
        )
        joined = gpd.sjoin_nearest(left, right, distance_col="distance")
        self.assertEqual(len(joined), 1)
        self.assertEqual(joined.iloc[0]["stop"], "near")
        self.assertAlmostEqual(joined.iloc[0]["distance"], 5.0)

    def test_union_all(self):
        import geopandas as gpd
        from shapely.geometry import box

        shapes = gpd.GeoSeries([box(0, 0, 2, 2), box(1, 0, 3, 2)], crs=3857)
        union = shapes.union_all()
        self.assertTrue(union.equals(box(0, 0, 3, 2)))
        self.assertAlmostEqual(union.area, 6.0)

    def test_groupby_apply(self):
        import pandas as pd

        frame = pd.DataFrame({"group": ["a", "a", "b"], "value": [2, 3, 7]})

        def total(group):
            self.assertNotIn("group", group.columns)
            return group["value"].sum()

        result = frame.groupby("group").apply(total, include_groups=False)
        self.assertEqual(result.to_dict(), {"a": 5, "b": 7})

    def test_spearman(self):
        import pandas as pd

        # Rank differences 0, -1, 1, 0 give rho = 1 - 6*2/(4*(16-1)).
        rho = pd.Series([1, 2, 3, 4]).corr(pd.Series([1, 3, 2, 4]), method="spearman")
        self.assertAlmostEqual(rho, 0.8)

    def test_matplotlib_png(self):
        import matplotlib

        matplotlib.use("Agg", force=True)
        import matplotlib.pyplot as plt

        self.assertEqual(matplotlib.get_backend().lower(), "agg")
        path = self.root / "plot.png"
        fig, ax = plt.subplots()
        try:
            ax.plot([0, 1], [0, 1])
            fig.savefig(path)
        finally:
            plt.close(fig)
        self.assertGreater(path.stat().st_size, 8)
        self.assertEqual(path.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")

    @unittest.skipUnless(importlib.util.find_spec("docx"), "optional python-docx unavailable")
    def test_docx(self):
        from docx import Document

        path = self.root / "sample.docx"
        document = Document()
        document.add_paragraph("Synthetic transport check")
        document.save(path)
        reopened = Document(path)
        self.assertEqual([p.text for p in reopened.paragraphs], ["Synthetic transport check"])


if __name__ == "__main__":
    # Set before importing Matplotlib. Run outside the working directory so a
    # local matplotlibrc cannot override defaults. All artifacts are temporary.
    with tempfile.TemporaryDirectory(prefix="transport-mpl-") as directory:
        previous_directory = Path.cwd()
        previous_env = {
            key: os.environ.get(key)
            for key in ("MPLCONFIGDIR", "XDG_CACHE_HOME", "MATPLOTLIBRC")
        }
        try:
            os.environ["MPLCONFIGDIR"] = str(Path(directory) / "config")
            os.environ["XDG_CACHE_HOME"] = str(Path(directory) / "cache")
            os.environ["MATPLOTLIBRC"] = directory
            os.chdir(directory)
            unittest.main(verbosity=2)
        finally:
            os.chdir(previous_directory)
            for key, value in previous_env.items():
                if value is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = value
